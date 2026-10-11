import csv
import io
from datetime import date, datetime
from pathlib import Path

from django.contrib.auth import get_user_model
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.utils.text import slugify
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .models import SchoolClass, SchoolMembership, Student, Subject, Teacher
from .models_onboarding import OnboardingRecord


User = get_user_model()
MAX_IMPORT_ROWS = 500

STUDENT_COLUMNS = [
    'first_name',
    'last_name',
    'class_name',
    'class_stream',
    'date_of_birth',
    'gender',
    'guardian_name',
    'guardian_email',
    'guardian_phone',
]
TEACHER_COLUMNS = [
    'first_name',
    'last_name',
    'email',
    'phone',
    'subjects',
    'classes',
]

TEMPLATE_EXAMPLES = {
    OnboardingRecord.PERSON_STUDENT: [
        'Amara', 'Okafor', 'JSS 1', 'A', '2013-04-18', 'F',
        'Chika Okafor', 'chika@example.com', '08030000000',
    ],
    OnboardingRecord.PERSON_TEACHER: [
        'Musa', 'Bello', 'musa@example.com', '08030000000',
        'Mathematics; Basic Science', 'JSS 1 A; JSS 2 A',
    ],
}


def columns_for(person_type):
    if person_type == OnboardingRecord.PERSON_TEACHER:
        return TEACHER_COLUMNS
    return STUDENT_COLUMNS


def template_response(person_type, file_format):
    columns = columns_for(person_type)
    label = 'teachers' if person_type == OnboardingRecord.PERSON_TEACHER else 'students'
    if file_format == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{label}_onboarding_template.csv"'
        response.write('\ufeff')
        writer = csv.writer(response)
        writer.writerow(columns)
        return response

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Onboarding'
    header_fill = PatternFill('solid', fgColor='DDE8FF')
    for column_number, column_name in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=column_number, value=column_name)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        sheet.column_dimensions[cell.column_letter].width = max(16, len(column_name) + 3)
    for column_number, value in enumerate(TEMPLATE_EXAMPLES[person_type], start=1):
        sheet.cell(row=2, column=column_number, value=value)
    sheet.freeze_panes = 'A2'
    instructions = workbook.create_sheet('Instructions')
    instructions.append(['How to use this template'])
    instructions['A1'].font = Font(bold=True)
    instructions.append(['Replace the example row with your records. Keep the column headings unchanged.'])
    instructions.append(['Separate multiple subjects or classes with a semicolon.'])
    instructions.append(['Dates use YYYY-MM-DD. Gender accepts M, F, Male, or Female.'])
    instructions.column_dimensions['A'].width = 92
    stream = io.BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{label}_onboarding_template.xlsx"'
    return response


def parse_spreadsheet(upload, person_type):
    extension = Path(upload.name).suffix.lower()
    columns = columns_for(person_type)
    if extension == '.csv':
        text = upload.read().decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(text))
        headers = [normalize_header(item) for item in (reader.fieldnames or [])]
        missing = [column for column in columns if column not in headers]
        if missing:
            raise ValueError(f"Missing columns: {', '.join(missing)}.")
        raw_rows = []
        for raw in reader:
            normalized = {normalize_header(key): clean_cell(value) for key, value in raw.items() if key}
            if any(normalized.values()):
                raw_rows.append({column: normalized.get(column, '') for column in columns})
    elif extension == '.xlsx':
        workbook = load_workbook(upload, read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        try:
            headers = [normalize_header(item) for item in next(values)]
        except StopIteration as exc:
            raise ValueError('The spreadsheet is empty.') from exc
        missing = [column for column in columns if column not in headers]
        if missing:
            raise ValueError(f"Missing columns: {', '.join(missing)}.")
        header_map = {name: index for index, name in enumerate(headers)}
        raw_rows = []
        for values_row in values:
            row = {
                column: clean_cell(values_row[header_map[column]] if header_map[column] < len(values_row) else '')
                for column in columns
            }
            if any(row.values()):
                raw_rows.append(row)
    else:
        raise ValueError('Upload a CSV or XLSX spreadsheet.')

    if not raw_rows:
        raise ValueError('The spreadsheet does not contain any records.')
    if len(raw_rows) > MAX_IMPORT_ROWS:
        raise ValueError(f'Import no more than {MAX_IMPORT_ROWS} records at a time.')
    return raw_rows


def rows_from_post(post, person_type):
    try:
        row_count = min(int(post.get('row_count', 0)), MAX_IMPORT_ROWS)
    except (TypeError, ValueError):
        row_count = 0
    columns = columns_for(person_type)
    return [
        {column: clean_cell(post.get(f'row_{row_index}_{column}', '')) for column in columns}
        for row_index in range(row_count)
    ]


def validate_rows(school, person_type, rows):
    reviewed = []
    seen_people = set()
    for index, raw_row in enumerate(rows, start=1):
        values = {column: clean_cell(raw_row.get(column, '')) for column in columns_for(person_type)}
        errors = {}
        for field_name in ['first_name', 'last_name']:
            if not values.get(field_name):
                errors[field_name] = 'Required.'

        if person_type == OnboardingRecord.PERSON_STUDENT:
            _validate_student_row(school, values, errors)
        else:
            _validate_teacher_row(school, values, errors)

        duplicate_key = (
            values.get('first_name', '').lower(),
            values.get('last_name', '').lower(),
            values.get('date_of_birth', '') if person_type == OnboardingRecord.PERSON_STUDENT else values.get('email', '').lower(),
        )
        if duplicate_key in seen_people:
            errors['first_name'] = 'This person appears more than once in the spreadsheet.'
        seen_people.add(duplicate_key)
        reviewed.append({
            'number': index,
            'values': values,
            'errors': errors,
            'valid': not errors,
        })
    return reviewed


def _validate_student_row(school, values, errors):
    if not values.get('class_name'):
        errors['class_name'] = 'Required.'
    school_class = find_school_class(school, values.get('class_name'), values.get('class_stream'))
    if values.get('class_name') and not school_class:
        errors['class_name'] = 'No matching class was found. Use the class name from Academic Setup.'

    parsed_date = parse_date(values.get('date_of_birth'))
    if not parsed_date:
        errors['date_of_birth'] = 'Use a valid date in YYYY-MM-DD format.'
    else:
        values['date_of_birth'] = parsed_date.isoformat()

    gender = normalize_gender(values.get('gender'))
    if not gender:
        errors['gender'] = 'Use M, F, Male, or Female.'
    else:
        values['gender'] = gender

    if not (values.get('guardian_email') or values.get('guardian_phone')):
        errors['guardian_email'] = 'Add a guardian email or phone number for the invitation.'

    if school_class and parsed_date and values.get('first_name') and values.get('last_name'):
        duplicate = Student.all_objects.filter(
            school=school,
            school_class=school_class,
            date_of_birth=parsed_date,
            user__first_name__iexact=values['first_name'],
            user__last_name__iexact=values['last_name'],
        ).exists()
        if duplicate:
            errors['first_name'] = 'A matching student already exists in this class.'


def _validate_teacher_row(school, values, errors):
    if not (values.get('email') or values.get('phone')):
        errors['email'] = 'Add an email or phone number for the invitation.'
    subject_names = split_list(values.get('subjects'))
    class_names = split_list(values.get('classes'))
    missing_subjects = [name for name in subject_names if not find_subject(school, name)]
    missing_classes = [name for name in class_names if not find_school_class_label(school, name)]
    if missing_subjects:
        errors['subjects'] = f"Not found: {', '.join(missing_subjects)}."
    if missing_classes:
        errors['classes'] = f"Not found: {', '.join(missing_classes)}."
    if values.get('email') and Teacher.all_objects.filter(
        school=school,
        user__email__iexact=values['email'],
    ).exists():
        errors['email'] = 'A teacher with this email already exists.'


@transaction.atomic
def import_rows(school, person_type, reviewed_rows, created_by=None):
    records = []
    for reviewed in reviewed_rows:
        if not reviewed.get('valid'):
            continue
        records.append(create_basic_person(
            school,
            person_type,
            reviewed['values'],
            created_by=created_by,
            source=OnboardingRecord.SOURCE_BULK,
        ))
    return records


@transaction.atomic
def create_basic_person(school, person_type, values, created_by=None, source=OnboardingRecord.SOURCE_MANUAL):
    first_name = clean_cell(values.get('first_name'))
    last_name = clean_cell(values.get('last_name'))
    username = generate_username(school, first_name, last_name)
    email = clean_cell(values.get('email'))
    phone = clean_cell(values.get('phone'))

    if person_type == OnboardingRecord.PERSON_STUDENT:
        email = ''
        phone = clean_cell(values.get('guardian_phone'))
        role = 'student'
    else:
        role = 'subject_teacher'

    user = User(
        username=username,
        first_name=first_name,
        last_name=last_name,
        email=email,
        phone=phone,
        role=role,
        is_active=True,
    )
    user.set_unusable_password()
    user.save()
    SchoolMembership.objects.create(
        school=school,
        user=user,
        primary_role=role,
        is_active=True,
        portal_access_enabled=False,
        result_access_enabled=True,
    )

    if person_type == OnboardingRecord.PERSON_STUDENT:
        school_class = values.get('school_class')
        if not isinstance(school_class, SchoolClass):
            school_class = find_school_class(school, values.get('class_name'), values.get('class_stream'))
        birth_date = values.get('date_of_birth')
        if not isinstance(birth_date, date):
            birth_date = parse_date(birth_date)
        student = Student.all_objects.create(
            school=school,
            user=user,
            school_class=school_class,
            date_of_birth=birth_date,
            gender=normalize_gender(values.get('gender')),
            status='active',
        )
        return OnboardingRecord.all_objects.create(
            school=school,
            person_type=person_type,
            source=source,
            user=user,
            student=student,
            recipient_name=clean_cell(values.get('guardian_name')),
            recipient_email=clean_cell(values.get('guardian_email')),
            recipient_phone=clean_cell(values.get('guardian_phone')),
            created_by=created_by,
        )

    teacher = Teacher.all_objects.create(school=school, user=user, employee_id='')
    subject_values = values.get('subjects')
    if hasattr(subject_values, 'all'):
        teacher.subjects.set(subject_values)
    elif isinstance(subject_values, (list, tuple, set)) and all(hasattr(item, 'pk') for item in subject_values):
        teacher.subjects.set(subject_values)
    else:
        teacher.subjects.set(filter(None, [find_subject(school, name) for name in split_list(subject_values)]))
    class_values = values.get('classes')
    if hasattr(class_values, 'all'):
        teacher.classes.set(class_values)
    elif isinstance(class_values, (list, tuple, set)) and all(hasattr(item, 'pk') for item in class_values):
        teacher.classes.set(class_values)
    else:
        teacher.classes.set(filter(None, [find_school_class_label(school, name) for name in split_list(class_values)]))
    return OnboardingRecord.all_objects.create(
        school=school,
        person_type=person_type,
        source=source,
        user=user,
        teacher=teacher,
        recipient_name=user.get_full_name(),
        recipient_email=user.email,
        recipient_phone=user.phone,
        created_by=created_by,
    )


def ensure_onboarding_records(school, created_by=None):
    from .staffing import active_staff, active_students
    existing_students = set(OnboardingRecord.all_objects.filter(
        school=school, student__isnull=False
    ).values_list('student_id', flat=True))
    existing_teachers = set(OnboardingRecord.all_objects.filter(
        school=school, teacher__isnull=False
    ).values_list('teacher_id', flat=True))
    new_records = []
    for student in active_students(school).select_related('user'):
        if student.id not in existing_students:
            new_records.append(OnboardingRecord(
                school=school,
                person_type=OnboardingRecord.PERSON_STUDENT,
                source=OnboardingRecord.SOURCE_EXISTING,
                user=student.user,
                student=student,
                recipient_name=student.user.get_full_name(),
                recipient_email=student.user.email,
                recipient_phone=student.user.phone,
                created_by=created_by,
            ))
    for teacher in active_staff(school).select_related('user'):
        if teacher.id not in existing_teachers:
            new_records.append(OnboardingRecord(
                school=school,
                person_type=OnboardingRecord.PERSON_TEACHER,
                source=OnboardingRecord.SOURCE_EXISTING,
                user=teacher.user,
                teacher=teacher,
                recipient_name=teacher.user.get_full_name(),
                recipient_email=teacher.user.email,
                recipient_phone=teacher.user.phone,
                created_by=created_by,
            ))
    if new_records:
        OnboardingRecord.all_objects.bulk_create(new_records, ignore_conflicts=True)
    records = OnboardingRecord.all_objects.filter(school=school).select_related(
        'user', 'student__school_class', 'teacher'
    )
    for record in records:
        record.sync_progress()


def generate_username(school, first_name, last_name):
    name_part = slugify(f'{first_name}.{last_name}').replace('-', '.') or 'portal.user'
    school_part = slugify(school.slug).replace('-', '.')[:24]
    base = f'{name_part}.{school_part}'[:140].strip('.')
    candidate = base
    suffix = 2
    while User.objects.filter(username__iexact=candidate).exists():
        candidate = f'{base[:145 - len(str(suffix))]}.{suffix}'
        suffix += 1
    return candidate


def find_school_class(school, name, stream=''):
    name = clean_cell(name)
    stream = clean_cell(stream)
    if not name:
        return None
    queryset = SchoolClass.all_objects.filter(school=school, name__iexact=name)
    if stream:
        queryset = queryset.filter(stream__iexact=stream)
    return queryset.order_by('stream').first()


def find_school_class_label(school, label):
    label = clean_cell(label)
    if not label:
        return None
    exact = SchoolClass.all_objects.filter(school=school).filter(
        Q(name__iexact=label) | Q(stream__iexact=label)
    ).first()
    if exact:
        return exact
    for school_class in SchoolClass.all_objects.filter(school=school).order_by('name', 'stream'):
        if str(school_class).lower() == label.lower():
            return school_class
    return None


def find_subject(school, label):
    label = clean_cell(label)
    if not label:
        return None
    return Subject.all_objects.filter(school=school).filter(
        Q(name__iexact=label) | Q(code__iexact=label)
    ).first()


def split_list(value):
    if not value:
        return []
    return [item.strip() for item in str(value).replace(',', ';').split(';') if item.strip()]


def normalize_header(value):
    return slugify(clean_cell(value)).replace('-', '_')


def clean_cell(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = clean_cell(value)
    for date_format in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(value, date_format).date()
        except (TypeError, ValueError):
            continue
    return None


def normalize_gender(value):
    value = clean_cell(value).lower()
    if value in {'m', 'male'}:
        return 'M'
    if value in {'f', 'female'}:
        return 'F'
    return ''
